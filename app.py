import csv
import io
from datetime import datetime
from typing import Dict, Any

from flask import (
    Flask,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)

from ai_parser import parse_expense_text
from config import MONTHLY_BUDGET, SECRET_KEY
from database import get_db_connection, init_db
from expense_service import (
    add_expense as save_expense,
    delete_expense,
    get_dashboard_data,
    get_expense_by_id,
    get_expenses,
    update_expense,
)
from whisper_service import transcribe_audio
from ai_chat import generate_finance_response, get_financial_suggestions

app = Flask(__name__)
app.config["SECRET_KEY"] = SECRET_KEY
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024

init_db()


def _get_first_present_value(row, *keys):
    for key in keys:
        if isinstance(row, dict):
            value = row.get(key)
            if value not in (None, ""):
                return value
    return None


def _iter_uploaded_rows(uploaded):
    filename = (uploaded.filename or "").lower()
    content_bytes = uploaded.read()

    if filename.endswith(".csv"):
        content = content_bytes.decode("utf-8")
        return list(csv.DictReader(io.StringIO(content)))

    if filename.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required for Excel uploads") from exc

        workbook = load_workbook(io.BytesIO(content_bytes), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = []
        for header in rows[0]:
            headers.append(str(header).strip().lower() if header is not None else "")

        parsed_rows = []
        for row in rows[1:]:
            record = {}
            for index, header in enumerate(headers):
                if index < len(row):
                    record[header] = row[index]
                else:
                    record[header] = None
            parsed_rows.append(record)
        return parsed_rows

    if filename.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("pypdf is required for PDF uploads") from exc

        reader = PdfReader(io.BytesIO(content_bytes))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return [{"raw_text": text}]

    raise ValueError("Unsupported file type. Please upload a CSV, .xlsx, or PDF file.")


def get_user_by_email(email: str):
    conn = get_db_connection()
    try:
        row = conn.execute(
            "SELECT id, name, email, password FROM users WHERE email = ?",
            (email,),
        ).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def create_user(name: str, email: str, password: str) -> bool:
    conn = get_db_connection()
    try:
        conn.execute(
            "INSERT INTO users (name, email, password, created_at) VALUES (?, ?, ?, ?)",
            (name, email, password, datetime.now().isoformat()),
        )
        conn.commit()
        return True
    except Exception:
        conn.rollback()
        return False
    finally:
        conn.close()


@app.route("/")
def home():
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")

        user = get_user_by_email(email)
        if user and user["password"] == password:
            session["user"] = email
            flash("Login Successful!", "success")
            return redirect(url_for("dashboard"))

        flash("Invalid Email or Password", "danger")

    return render_template("login.html")


@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "POST":
        name = request.form.get("fullname", "").strip()
        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")

        if password != confirm:
            flash("Passwords do not match.", "danger")
            return redirect(url_for("signup"))

        if get_user_by_email(email):
            flash("Email already exists.", "warning")
            return redirect(url_for("signup"))

        if create_user(name, email, password):
            flash("Account created successfully.", "success")
            return redirect(url_for("login"))

        flash("Unable to create account.", "danger")
        return redirect(url_for("signup"))

    return render_template("signup.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Logged out successfully.", "success")
    return redirect(url_for("login"))


def login_required():
    return "user" in session


def build_page_data(email: str) -> Dict[str, Any]:
    data = get_dashboard_data(email)
    user = get_user_by_email(email)
    data["user"] = {"name": user["name"] if user else email, "email": email}
    return data


@app.route("/dashboard")
def dashboard():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("dashboard.html", **build_page_data(session["user"]))


@app.route("/expenses")
def expenses():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("expenses.html", **build_page_data(session["user"]))


@app.route("/add_expense", methods=["GET", "POST"])
def add_expense():
    if not login_required():
        return redirect(url_for("login"))

    if request.method == "POST":
        data = {
            "merchant": request.form.get("merchant") or "Unknown Merchant",
            "category": request.form.get("category") or "Miscellaneous",
            "amount": request.form.get("amount") or 0,
            "payment_method": request.form.get("payment") or "Cash",
            "date": request.form.get("date") or datetime.now().strftime("%Y-%m-%d"),
            "description": request.form.get("notes") or "Manual entry",
        }
        result = save_expense(session["user"], data)
        if result["success"]:
            flash("Expense saved successfully.", "success")
        else:
            flash("Unable to save expense.", "danger")
        return redirect(url_for("expenses"))

    return render_template("add_expense.html", **build_page_data(session["user"]))


@app.route("/voice")
def voice():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("voice.html", **build_page_data(session["user"]))


@app.route("/upload")
def upload():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("upload.html", **build_page_data(session["user"]))


@app.route("/analytics")
def analytics():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("analytics.html", **build_page_data(session["user"]))


@app.route("/comparison")
def comparison():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("comparison.html", **build_page_data(session["user"]))


@app.route("/categories")
def categories():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("categories.html", **build_page_data(session["user"]))


@app.route("/reports")
def reports():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("reports.html", **build_page_data(session["user"]))


@app.route("/budget")
def budget():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("budget.html", **build_page_data(session["user"]))


@app.route("/suggestions")
def suggestions():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("suggestions.html", **build_page_data(session["user"]))


@app.route("/settings")
def settings():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("settings.html", **build_page_data(session["user"]))


@app.route("/thomas-ai")
def thomas_ai():
    if not login_required():
        return redirect(url_for("login"))
    return render_template("thomas_ai.html", **build_page_data(session["user"]))


@app.route("/api/thomas-ai", methods=["POST"])
def api_thomas_ai():
    if not login_required():
        return jsonify({"success": False, "message": "Please login first."}), 401

    data = request.get_json(silent=True) or {}
    user_message = data.get("message", "").strip()
    
    if not user_message:
        return jsonify({"success": False, "message": "Message is required."}), 400
    
    try:
        response = generate_finance_response(user_message)
        return jsonify({
            "success": True,
            "response": response,
            "sender": "Thomas Cook AI"
        })
    except Exception as exc:
        return jsonify({
            "success": False,
            "message": f"Error: {str(exc)}"
        }), 500


@app.route("/api/thomas-ai/suggestions", methods=["GET"])
def api_thomas_ai_suggestions():
    if not login_required():
        return jsonify({"success": False}), 401
    
    try:
        expenses = get_expenses(session["user"])
        budget = session.get("monthly_budget", MONTHLY_BUDGET)
        suggestions = get_financial_suggestions(expenses, budget)
        return jsonify({
            "success": True,
            "suggestions": suggestions
        })
    except Exception as exc:
        return jsonify({
            "success": False,
            "message": str(exc)
        }), 500


@app.route("/ai_expense", methods=["POST"])
def ai_expense():
    if not login_required():
        return jsonify({"success": False, "message": "Please login first."}), 401

    if request.is_json:
        payload = request.get_json(silent=True) or {}
        text = payload.get("expense_text") or ""
    else:
        text = request.form.get("expense_text") or ""

    try:
        parsed = parse_expense_text(text)
        result = save_expense(session["user"], {**parsed, "description": text})
        return jsonify(
            {
                "success": True,
                "amount": parsed["amount"],
                "category": parsed["category"],
                "merchant": parsed["merchant"],
                "date": parsed["date"],
                "expense": result["expense"],
            }
        )
    except ValueError as exc:
        return jsonify({"success": False, "message": str(exc)}), 400


@app.route("/voice/upload", methods=["POST"])
def voice_upload():
    if not login_required():
        return jsonify({"success": False, "message": "Please login first."}), 401

    audio_file = request.files.get("audio")
    if not audio_file:
        return jsonify({"success": False, "message": "No audio provided."}), 400

    try:
        transcription = transcribe_audio(audio_file)
    except Exception as exc:
        return jsonify({"success": False, "message": str(exc)}), 400

    parsed = parse_expense_text(transcription)
    save_expense(session["user"], {**parsed, "description": transcription})
    return jsonify(
        {
            "success": True,
            "transcription": transcription,
            "amount": parsed["amount"],
            "category": parsed["category"],
            "merchant": parsed["merchant"],
            "date": parsed["date"],
        }
    )


@app.route("/upload", methods=["POST"])
def upload_expenses():
    if not login_required():
        return jsonify({"success": False, "message": "Please login first."}), 401

    uploaded = request.files.get("expenseFile")
    if not uploaded:
        return jsonify({"success": False, "message": "No file selected."}), 400

    imported = 0
    skipped = 0
    try:
        rows = _iter_uploaded_rows(uploaded)

        # If PDF was uploaded, _iter_uploaded_rows returns a single dict with raw_text
        if rows and isinstance(rows[0], dict) and "raw_text" in rows[0]:
            for row in rows:
                text = row.get("raw_text", "") or ""
                try:
                    parsed = parse_expense_text(text)
                except Exception:
                    skipped += 1
                    continue

                save_expense(
                    session["user"],
                    {
                        "merchant": parsed.get("merchant", "Imported Merchant"),
                        "category": parsed.get("category", "Miscellaneous"),
                        "amount": parsed.get("amount"),
                        "payment_method": parsed.get("payment_method", "Card"),
                        "date": parsed.get("date", datetime.now().strftime("%Y-%m-%d")),
                        "description": text,
                    },
                )
                imported += 1

        else:
            for row in rows:
                amount = _get_first_present_value(row, "amount", "total", "value")
                if amount in (None, ""):
                    skipped += 1
                    continue

                merchant = _get_first_present_value(row, "merchant", "vendor", "name", "store") or "Imported Merchant"
                category = _get_first_present_value(row, "category", "type", "expense_category") or "Miscellaneous"
                payment_method = _get_first_present_value(row, "payment_method", "method", "payment", "mode") or "Card"
                date_value = _get_first_present_value(row, "date", "spent_on", "transaction_date") or datetime.now().strftime("%Y-%m-%d")

                save_expense(
                    session["user"],
                    {
                        "merchant": merchant,
                        "category": category,
                        "amount": amount,
                        "payment_method": payment_method,
                        "date": date_value,
                        "description": "Imported from upload",
                    },
                )
                imported += 1
    except Exception as exc:
        return jsonify({"success": False, "message": f"Could not read the uploaded file: {exc}"}), 400

    return jsonify(
        {
            "success": True,
            "imported": imported,
            "skipped": skipped,
            "message": f"Imported {imported} expense(s) successfully.",
        }
    )


@app.route("/edit/<int:expense_id>", methods=["GET", "POST"])
def edit_expense(expense_id: int):
    if not login_required():
        return redirect(url_for("login"))

    user_email = session["user"]
    expense = get_expense_by_id(user_email, expense_id)
    if not expense:
        flash("Expense not found.", "warning")
        return redirect(url_for("expenses"))

    if request.method == "POST":
        data = {
            "merchant": request.form.get("merchant") or "Unknown Merchant",
            "category": request.form.get("category") or "Miscellaneous",
            "amount": request.form.get("amount") or 0,
            "payment_method": request.form.get("payment") or "Cash",
            "date": request.form.get("date") or datetime.now().strftime("%Y-%m-%d"),
            "description": request.form.get("notes") or "Manual entry",
        }
        if update_expense(user_email, expense_id, data):
            flash("Expense updated successfully.", "success")
        else:
            flash("Unable to update expense.", "danger")
        return redirect(url_for("expenses"))

    return render_template(
        "edit_expense.html",
        expense=expense,
        **build_page_data(user_email),
    )


@app.route("/delete/<int:expense_id>")
def delete_expense_route(expense_id: int):
    if not login_required():
        return redirect(url_for("login"))

    user_email = session["user"]
    if delete_expense(user_email, expense_id):
        flash("Expense deleted successfully.", "success")
    else:
        flash("Could not delete expense.", "danger")
    return redirect(url_for("expenses"))


@app.route("/export/csv")
def export_csv():
    if not login_required():
        return redirect(url_for("login"))

    expenses = get_expenses(session["user"])
    csv_file = io.StringIO()
    writer = csv.writer(csv_file)
    writer.writerow(["Date", "Merchant", "Category", "Amount", "Payment Method", "Description"])
    for expense in expenses:
        writer.writerow(
            [
                expense["date"],
                expense["merchant"],
                expense["category"],
                expense["amount"],
                expense.get("payment_method", ""),
                expense.get("description", ""),
            ]
        )

    output = io.BytesIO(csv_file.getvalue().encode("utf-8"))
    output.seek(0)
    return send_file(
        output,
        mimetype="text/csv",
        as_attachment=True,
        download_name="expenses.csv",
    )


@app.route("/export/excel")
def export_excel():
    if not login_required():
        return redirect(url_for("login"))

    expenses = get_expenses(session["user"])
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Expenses"
    headers = ["Date", "Merchant", "Category", "Amount", "Payment Method", "Description"]
    sheet.append(headers)
    for expense in expenses:
        sheet.append(
            [
                expense["date"],
                expense["merchant"],
                expense["category"],
                expense["amount"],
                expense.get("payment_method", ""),
                expense.get("description", ""),
            ]
        )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="expenses.xlsx",
    )


@app.route("/export/pdf")
def export_pdf():
    if not login_required():
        return redirect(url_for("login"))

    expenses = get_expenses(session["user"])
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen.canvas import Canvas
    except ImportError:
        flash("PDF export requires the reportlab package.", "danger")
        return redirect(url_for("reports"))

    output = io.BytesIO()
    canvas = Canvas(output, pagesize=letter)
    width, height = letter
    y = height - 40
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(40, y, "Expense Report")
    y -= 30
    canvas.setFont("Helvetica-Bold", 10)
    headers = ["Date", "Merchant", "Category", "Amount", "Payment", "Description"]
    positions = [40, 110, 240, 320, 390, 470]
    for index, header in enumerate(headers):
        canvas.drawString(positions[index], y, header)
    y -= 16
    canvas.setFont("Helvetica", 9)
    for expense in expenses:
        if y < 80:
            canvas.showPage()
            y = height - 40
            canvas.setFont("Helvetica-Bold", 10)
            for index, header in enumerate(headers):
                canvas.drawString(positions[index], y, header)
            y -= 16
            canvas.setFont("Helvetica", 9)

        values = [
            expense["date"],
            expense["merchant"],
            expense["category"],
            str(expense["amount"]),
            expense.get("payment_method", ""),
            expense.get("description", ""),
        ]
        for index, value in enumerate(values):
            canvas.drawString(positions[index], y, str(value)[:25])
        y -= 14

    canvas.save()
    output.seek(0)
    return send_file(
        output,
        mimetype="application/pdf",
        as_attachment=True,
        download_name="expenses.pdf",
    )


@app.route("/generate-report", methods=["POST"])
def generate_report():
    if not login_required():
        return redirect(url_for("login"))

    user_email = session["user"]
    from_date = request.form.get("from_date")
    to_date = request.form.get("to_date")
    reports = get_expenses(user_email)

    if from_date:
        reports = [expense for expense in reports if expense["date"] >= from_date]
    if to_date:
        reports = [expense for expense in reports if expense["date"] <= to_date]

    data = build_page_data(user_email)
    data["reports"] = reports
    data["report_type"] = request.form.get("report_type", "Monthly")
    return render_template("reports.html", **data)


@app.route("/set-budget", methods=["POST"])
def set_budget():
    if not login_required():
        return redirect(url_for("login"))

    budget_value = request.form.get("monthly_budget", MONTHLY_BUDGET)
    try:
        session["monthly_budget"] = float(budget_value)
    except ValueError:
        session["monthly_budget"] = MONTHLY_BUDGET
    flash("Budget updated successfully.", "success")
    return redirect(url_for("budget"))


@app.route("/api/analytics")
def api_analytics():
    if not login_required():
        return jsonify({"success": False}), 401

    data = build_page_data(session["user"])
    return jsonify(
        {
            "success": True,
            "total_expense": data["total_expense"],
            "average_daily": data["avg_daily"],
            "highest_category": data["highest_category"],
            "predicted_expense": data["predicted_expense"],
            "insights": data.get("insights", data.get("alerts", [])),
            "months": [datetime.now().strftime("%b")],
            "expenses": [data["monthly_total"]],
            "categories": [item["name"] for item in data.get("categories", [])],
            "category_amounts": [item["amount"] for item in data.get("categories", [])],
            "future_months": ["Next Month"],
            "future_prediction": [data["predicted_expense"]],
        }
    )


@app.route("/api/budget")
def api_budget():
    if not login_required():
        return jsonify({"success": False}), 401

    data = build_page_data(session["user"])
    return jsonify(
        {
            "success": True,
            "monthly_budget": data.get("monthly_budget", MONTHLY_BUDGET),
            "total_spent": data.get("total_spent", data.get("monthly_total", 0)),
            "remaining_budget": data.get("remaining_budget", data.get("monthly_budget", MONTHLY_BUDGET) - data.get("monthly_total", 0)),
            "usage_percentage": data.get("usage_percentage", 0),
            "categories": [
                {
                    "name": item["category"],
                    "limit": item["limit"],
                    "spent": item["spent"],
                    "remaining": item["remaining"],
                }
                for item in data.get("category_budget", [])
            ],
            "alerts": data.get("alerts", []),
            "recommendations": data.get("recommendations", []),
        }
    )


@app.route("/api/comparison")
def api_comparison():
    if not login_required():
        return jsonify({"success": False}), 401

    data = build_page_data(session["user"])
    return jsonify(
        {
            "success": True,
            "current_month": data.get("current_month", 0),
            "previous_month": data.get("previous_month", 0),
            "difference": data.get("difference", 0),
            "percentage": data.get("percentage", 0),
            "categories": data.get("comparison_data", []),
            "insights": data.get("comparison_insights", []),
            "months": ["Previous", "Current"],
            "current": [data.get("previous_month", 0), data.get("current_month", 0)],
            "previous": [data.get("previous_month", 0), data.get("previous_month", 0)],
            "growth": [0, data.get("percentage", 0)],
        }
    )


@app.route("/api/suggestions")
def api_suggestions():
    if not login_required():
        return jsonify({"success": False}), 401

    data = build_page_data(session["user"])
    return jsonify(
        {
            "success": True,
            "predicted_expense": data.get("predicted_expense", 0),
            "expected_savings": data.get("expected_savings", 0),
            "highest_category": data.get("highest_category", "None"),
            "budget_status": data.get("budget_status", "Healthy"),
            "recommendations": data.get("recommendations", []),
            "savings": data.get("savings", []),
            "alerts": data.get("alerts", []),
            "months": ["Current", "Next"],
            "prediction": [data.get("monthly_total", 0), data.get("predicted_expense", 0)],
        }
    )


if __name__ == "__main__":
    app.run(debug=True)